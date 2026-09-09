"""Estágio 0: qualquer formato vira texto. Sem IA e sem custo.

O arquivo original é preservado: é o que a interface mostra ao recrutador e é
o que sustenta o OCR quando o PDF é escaneado.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, Optional

from .config import MAX_CHARS_CURRICULO, MAX_PAGINAS_PDF, MIN_CHARS_TEXTO_UTIL

log = logging.getLogger(__name__)

# Assinaturas de arquivo. Extensão trocada é o erro de upload mais comum do RH.
ASSINATURAS = {
    ".pdf": (b"%PDF",),
    ".docx": (b"PK\x03\x04",),
    ".xlsx": (b"PK\x03\x04",),
    ".doc": (b"\xd0\xcf\x11\xe0", b"{\\rtf", b"PK\x03\x04"),
    ".rtf": (b"{\\rtf", b"\xd0\xcf\x11\xe0"),
}


@dataclass
class Documento:
    """Um currículo pronto para o parser."""
    candidato_id: str
    arquivo: str
    texto: str = ""
    origem: str = "arquivo"                # arquivo | planilha
    origem_hash: str = ""                  # hash do arquivo enviado (nome dele no disco)
    extensao: str = ""
    escaneado: bool = False                # PDF sem texto embutido: caso de OCR
    origem_texto: str = "direto"           # direto | ocr
    ocr_confianca: str = ""                # alta | baixa (vazio quando não houve OCR)
    estado_extracao: str = "pronto"        # pronto | pendente | insuficiente | falhou
    avisos: list[str] = field(default_factory=list)


def hash_conteudo(dados: bytes) -> str:
    return hashlib.sha256(dados).hexdigest()


def id_candidato(dados: bytes) -> str:
    """16 hex = 64 bits. Colisão exigiria bilhões de currículos na mesma conta."""
    return hash_conteudo(dados)[:16]


def _limpar(texto: str, avisos: Optional[list] = None) -> str:
    texto = texto.replace("\x00", " ").replace("﻿", "")   # nulo e BOM
    # PDF com fonte sem mapa de caracteres devolve marcador no lugar do glifo;
    # sem isto, "(cid:127)" entra onde deveria haver um marcador de lista.
    texto = re.sub(r"\(cid:\d+\)", "- ", texto)
    texto = re.sub(r"[ \t ]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    texto = texto.strip()
    if len(texto) > MAX_CHARS_CURRICULO:
        if avisos is not None:
            avisos.append(
                f"currículo muito longo ({len(texto)} caracteres); "
                f"avaliados os primeiros {MAX_CHARS_CURRICULO}"
            )
        texto = texto[:MAX_CHARS_CURRICULO]
    return texto


def conferir_assinatura(dados: bytes, extensao: str) -> Optional[str]:
    """Devolve o problema, ou None. Barra o executável renomeado para .pdf."""
    esperadas = ASSINATURAS.get(extensao)
    if not esperadas:
        return None
    cabeca = dados[:8]
    if any(cabeca.startswith(a) for a in esperadas):
        return None
    return f"o conteúdo não parece um arquivo {extensao}"


# ---------- PDF ----------

def _texto_pdf(dados: bytes, avisos: list[str]) -> tuple[str, int]:
    import pdfplumber

    partes: list[str] = []
    with pdfplumber.open(io.BytesIO(dados)) as pdf:
        total = len(pdf.pages)
        for pagina in pdf.pages[:MAX_PAGINAS_PDF]:
            partes.append(pagina.extract_text() or "")
            for tabela in pagina.extract_tables() or []:       # currículo em tabela é comum
                for linha in tabela:
                    celulas = [c for c in linha if c]
                    if celulas:
                        partes.append(" | ".join(celulas))
    if total > MAX_PAGINAS_PDF:
        avisos.append(f"{total} páginas; lidas as {MAX_PAGINAS_PDF} primeiras")
    return _limpar("\n".join(partes), avisos), total


# ---------- DOCX / DOC ----------

def _texto_docx(dados: bytes, avisos: list[str]) -> str:
    import docx

    doc = docx.Document(io.BytesIO(dados))
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:                                  # muito currículo é tabela
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    for secao in doc.sections:                                 # contato costuma ir no cabeçalho
        for parte in (secao.header, secao.footer):
            for p in getattr(parte, "paragraphs", []):
                if p.text.strip():
                    partes.append(p.text)
    return _limpar("\n".join(partes), avisos)


# O instalador do LibreOffice no Windows não põe o soffice no PATH, e no Linux
# alguns pacotes usam nomes diferentes. Procurar nos lugares óbvios evita recusar
# `.doc` de uma máquina que tem o programa instalado — que seria uma mensagem de
# erro mentindo para o recrutador.
_CAMINHOS_LIBREOFFICE = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/usr/lib/libreoffice/program/soffice",
    "/opt/libreoffice/program/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


def achar_libreoffice() -> Optional[str]:
    """O caminho do soffice, ou None. Primeiro o PATH, depois os lugares usuais."""
    for nome in ("soffice", "libreoffice", "soffice.exe"):
        achado = shutil.which(nome)
        if achado:
            return achado
    for caminho in _CAMINHOS_LIBREOFFICE:
        if Path(caminho).exists():
            return caminho
    return None


def _texto_doc_legado(dados: bytes, avisos: list[str]) -> str:
    """.doc e .rtf antigos: converte com LibreOffice, se existir na máquina."""
    binario = achar_libreoffice()
    if not binario:
        raise RuntimeError(
            "arquivo .doc/.rtf exige LibreOffice instalado no servidor "
            "(a imagem Docker do produto já vem com ele)"
        )
    with tempfile.TemporaryDirectory() as tmp:
        # A extensão precisa bater com o conteúdo: o LibreOffice escolhe o filtro
        # de importação por ela, e um `.rtf` entregue como `.doc` volta ilegível.
        extensao = ".rtf" if dados[:5].lstrip().startswith(b"{\\rtf") else ".doc"
        entrada = Path(tmp) / f"entrada{extensao}"
        entrada.write_bytes(dados)
        try:
            subprocess.run(
                [binario, "--headless", "--norestore", "--convert-to", "docx",
                 "--outdir", tmp, str(entrada)],
                capture_output=True, timeout=120, check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOffice travou ao converter o arquivo") from exc
        convertidos = list(Path(tmp).glob("*.docx"))
        if not convertidos:
            raise RuntimeError("LibreOffice não conseguiu converter o arquivo")
        return _texto_docx(convertidos[0].read_bytes(), avisos)


# ---------- Planilhas exportadas de ATS / LinkedIn ----------

COLUNAS_NOME = {"nome", "name", "candidato", "candidate", "full name",
                "nome completo", "nome do candidato"}
COLUNAS_IGNORADAS = {"id", "created at", "updated at", "url da vaga",
                     "data de inscrição", "origem", "status", "etapa", "fonte"}


def _decodificar(dados: bytes) -> str:
    for codificacao in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace")


def _linhas_planilha(dados: bytes, arquivo: str) -> Iterator[dict]:
    if arquivo.lower().endswith(".csv"):
        texto = _decodificar(dados)
        try:
            dialeto = csv.Sniffer().sniff(texto[:4000], delimiters=",;\t|")
        except csv.Error:
            dialeto = csv.excel
        yield from csv.DictReader(io.StringIO(texto), dialect=dialeto)
        return

    try:
        from openpyxl import load_workbook
    except ImportError as exc:                                 # pragma: no cover
        raise RuntimeError("leitura de .xlsx exige openpyxl instalado") from exc

    wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    try:
        for aba in wb.worksheets:
            linhas = aba.iter_rows(values_only=True)
            cabecalho = [str(c or "").strip() for c in next(linhas, ())]
            if not any(cabecalho):
                continue
            for linha in linhas:
                yield {k: v for k, v in zip(cabecalho, linha) if k}
    finally:
        wb.close()


def _documentos_de_planilha(dados: bytes, arquivo: str, origem_hash: str) -> list[Documento]:
    docs: list[Documento] = []
    for i, linha in enumerate(_linhas_planilha(dados, arquivo)):
        campos: list[str] = []
        nome_bruto = ""
        for chave, valor in linha.items():
            if valor is None or str(valor).strip() == "":
                continue
            chave_norm = str(chave).strip().lower()
            if chave_norm in COLUNAS_IGNORADAS:
                continue
            if chave_norm in COLUNAS_NOME:
                nome_bruto = str(valor).strip()
            campos.append(f"{str(chave).strip()}: {str(valor).strip()}")

        if len(campos) < 2:                                    # linha vazia ou só de controle
            continue

        avisos: list[str] = []
        texto = _limpar("\n".join(campos), avisos)
        rotulo = f"{arquivo} · linha {i + 2}" + (f" · {nome_bruto}" if nome_bruto else "")
        docs.append(Documento(
            candidato_id=id_candidato(f"{origem_hash}:{texto}".encode()),
            arquivo=rotulo,
            texto=texto,
            origem="planilha",
            origem_hash=origem_hash,
            extensao=Path(arquivo).suffix.lower(),
            avisos=avisos,
        ))
        _aplicar_piso_de_texto(docs[-1])
    return docs


# ---------- Entrada única ----------

def _aplicar_piso_de_texto(doc: Documento) -> None:
    """Piso de entrada: abaixo de MIN_CHARS_TEXTO_UTIL não há currículo ali.

    O mesmo limiar que o PDF já usava, agora valendo para DOCX, DOC, RTF, TXT e
    linha de planilha. No PDF ele significa "provavelmente escaneado, manda para
    o OCR"; aqui significa "o arquivo abriu, mas não tem texto que sustente uma
    avaliação" — e o destino é outro: cai no caminho de "não foi possível
    avaliar" que a tela, o CSV e o relatório já têm.

    Existe porque o piloto trouxe linhas de ATS com só nome, e-mail e cidade.
    Elas atravessavam como currículo, recebiam nota 0 e apareciam ao recrutador
    como quem foi medido e ficou por último — quando a verdade é que não havia o
    que medir.

    O texto extraído NÃO é apagado: a extração funcionou, e jogar fora o pouco
    que veio confundiria "não consegui ler o arquivo" com "li, e não há currículo
    aqui". O que muda é o estado, e com ele o destino.
    """
    if doc.estado_extracao == "pendente":     # o OCR ainda vai ler; outro caso
        return
    texto = (doc.texto or "").strip()
    if not texto or len(texto) >= MIN_CHARS_TEXTO_UTIL:
        return
    doc.estado_extracao = "insuficiente"
    doc.avisos.append(
        f"texto insuficiente para avaliar: {len(texto)} caracteres, "
        f"mínimo de {MIN_CHARS_TEXTO_UTIL}"
    )


def extrair(dados: bytes, arquivo: str) -> list[Documento]:
    """Recebe os bytes de um upload e devolve 1..N documentos prontos para o parser."""
    ext = Path(arquivo).suffix.lower()
    completo = hash_conteudo(dados)

    if ext in {".csv", ".xlsx"}:
        return _documentos_de_planilha(dados, arquivo, completo)

    doc = Documento(
        candidato_id=completo[:16],
        arquivo=arquivo,
        origem_hash=completo,
        extensao=ext,
    )
    try:
        if ext == ".pdf":
            doc.texto, _paginas = _texto_pdf(dados, doc.avisos)
            if len(doc.texto) < MIN_CHARS_TEXTO_UTIL:
                # PDF escaneado não é lido aqui. O OCR leva de 2 a 6 segundos por
                # página, e fazê-lo dentro da requisição de upload deixaria o
                # recrutador esperando minutos ao arrastar a pasta inteira — o
                # tempo suficiente para o proxy derrubar a conexão. Ele sai daqui
                # marcado, e a fila lê depois.
                doc.escaneado = True
                doc.estado_extracao = "pendente"
                doc.avisos.append("PDF escaneado; na fila para leitura por OCR")
        elif ext == ".docx":
            doc.texto = _texto_docx(dados, doc.avisos)
        elif ext in {".doc", ".rtf"}:
            doc.texto = _texto_doc_legado(dados, doc.avisos)
        elif ext in {".txt", ".md"}:
            doc.texto = _limpar(_decodificar(dados), doc.avisos)
        else:
            raise RuntimeError(f"formato não suportado: {ext or 'sem extensão'}")
    except Exception as exc:                                   # noqa: BLE001
        log.warning("falha ao extrair %s: %s", arquivo, exc)
        doc.avisos.append(f"falha na extração: {exc}")

    _aplicar_piso_de_texto(doc)
    return [doc]
